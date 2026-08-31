"""Convert the assembled Markdown report into a .docx laid out to the
department's UG project report format.

There is no pandoc on this machine, so the conversion is written directly against
python-docx. That turns out to be the right call anyway, because the format
specification is precise about things a generic converter gets wrong:

  General text        Times New Roman 12, 1.5 line spacing, justified
  Chapter name        Times New Roman 16, bold, centred
  Heading (x.y)       Times New Roman 14, bold, left aligned
  Sub-heading (x.y.z) Times New Roman 12, bold, left aligned
  Figure caption      Times New Roman 10, bold, centred, BELOW the figure
  Table caption       Times New Roman 10, bold, centred, ABOVE the table
  Abstract            Times New Roman 14, double spaced
  Bonafide            Times New Roman 14, double spaced
  Contents and lists  1.5 line spacing
  References          single spacing, left justified
  Front matter        lower-case roman page numbers; body restarts at arabic

Two structural details are handled rather than left to the typist. The
"Table of Contents" heading emits a real Word TOC *field* instead of a static
table -- because the headings below carry genuine Heading 1/2/3 styles, Word can
build and update the contents itself, and a hand-maintained list of page numbers
in a 120-page document is wrong the moment anything is edited. And the front
matter sits in its own section so it can carry roman numerals while the chapters
restart at 1.

Run:  python -m tools.make_docx   ->  writing/report/PROJECT_REPORT.docx
"""
from __future__ import annotations

import re
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

SRC = Path("writing/report/PROJECT_REPORT.md")
OUT = Path("writing/report/PROJECT_REPORT.docx")
PNG = Path("writing/figures/png")
DIAGRAM_PNG = Path("writing/figures/diagrams")

FONT = "Times New Roman"
MONO = "Consolas"
MAX_IMG_IN = 6.0          # printable width inside 1-inch margins on A4
MAX_IMG_H_IN = 8.2        # printable height, leaving room for the caption

#: Front-matter headings: each starts a new page and is centred at 16 pt.
FRONT = {"title page", "bonafide certificate", "acknowledgement",
         "table of contents", "list of figures", "list of tables",
         "abstract", "graphical abstract", "list of abbreviations",
         "abbreviations", "symbols"}

#: Typed at 14 pt double spacing, per the specification.
DOUBLE_14 = {"abstract", "bonafide certificate"}

#: Set 1.5 spacing and left alignment rather than justified.
LISTY = {"list of figures", "list of tables", "list of abbreviations",
         "abbreviations", "symbols"}


# --------------------------------------------------------------- helpers ----
def _a4(section) -> None:
    """A4 is 210 x 297 mm.  python-docx defaults to US Letter, and the
    department format specification asks for A4, so every section is set
    explicitly rather than left to the default."""
    section.page_width = Mm(210)
    section.page_height = Mm(297)


def _shade(par, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    par._p.get_or_add_pPr().append(shd)


def _page_number_format(section, fmt: str, start: int | None = None) -> None:
    # A section created with add_section() copies the previous section's
    # properties, pgNumType included. Appending a second one does not override
    # the first, so the inherited element is removed before the new one is added.
    for old in section._sectPr.findall(qn("w:pgNumType")):
        section._sectPr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    section._sectPr.append(pg)


def _add_page_field(par) -> None:
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for tag, attr, val in (("w:fldChar", "w:fldCharType", "begin"),
                           ("w:instrText", None, " PAGE "),
                           ("w:fldChar", "w:fldCharType", "end")):
        run = par.add_run()
        el = OxmlElement(tag)
        if attr:
            el.set(qn(attr), val)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = val
        run._r.append(el)
        run.font.name = FONT
        run.font.size = Pt(11)


def _add_toc_field(par) -> None:
    run = par.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    hint = OxmlElement("w:t")
    hint.text = "Right-click and choose Update Field to build the contents."
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, hint, end):
        run._r.append(el)


INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))", re.S)


def _runs(par, text: str, size: float = 12.0, bold_all: bool = False) -> None:
    text = text.replace("\\_", "_")
    text = re.sub(r"<sub>(.*?)</sub>", r"_\1", text)
    text = re.sub(r"<sup>(.*?)</sup>", r"^\1", text)
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = par.add_run(part[2:-2]); r.bold = True
        elif (part.startswith("*") and part.endswith("*") and len(part) > 2
              and not part.startswith("**")):
            r = par.add_run(part[1:-1]); r.italic = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.font.name = MONO
            r.font.size = Pt(size - 1.5)
            continue
        elif part.startswith("[") and "](" in part:
            r = par.add_run(part[1:part.index("](")])
        else:
            r = par.add_run(part)
        r.font.name = FONT
        r.font.size = Pt(size)
        if bold_all:
            r.bold = True


def _image_for(ref: str) -> Path | None:
    stem = Path(ref).stem
    for folder in (DIAGRAM_PNG, PNG):
        for ext in (".png", ".PNG", ".jpg", ".jpeg"):
            c = folder / f"{stem}{ext}"
            if c.exists():
                return c
    return None


def _fmt(par, *, spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         before=0, after=6) -> None:
    pf = par.paragraph_format
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    par.alignment = align


def _add_caption(doc, kind: str, title: str):
    """A real Word caption: literal label, SEQ field, then the title.

    Written as a field rather than as plain text so Word treats it exactly as
    References -> Insert Caption would: the numbers renumber themselves when
    anything moves, they appear in the Insert Cross-reference list, and
    Insert Table of Figures can collect them.
    """
    p = doc.add_paragraph(style="Caption")
    _fmt(p, spacing=1.0, align=WD_ALIGN_PARAGRAPH.CENTER, before=3, after=8)

    r = p.add_run(f"{kind} ")
    r.font.name = FONT; r.font.size = Pt(10); r.bold = True

    run = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = f" SEQ {kind} \* ARABIC "
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, end):
        run._r.append(el)
    run.font.name = FONT; run.font.size = Pt(10); run.bold = True

    r = p.add_run(f": {title}")
    r.font.name = FONT; r.font.size = Pt(10); r.bold = True
    return p


def _add_table_of(doc, kind: str):
    """An auto-generated List of Figures / List of Tables."""
    p = doc.add_paragraph()
    _fmt(p, spacing=1.5, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = f' TOC \h \z \c "{kind}" '
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    hint = OxmlElement("w:t")
    hint.text = f"Right-click and choose Update Field to list the {kind.lower()}s."
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, hint, end):
        run._r.append(el)
    return p


# ------------------------------------------------------------------ main ----
def build() -> Document:
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(6)

    cap = doc.styles["Caption"]
    cap.font.name = FONT
    cap.font.size = Pt(10)
    cap.font.bold = True
    cap.font.color.rgb = RGBColor(0, 0, 0)

    for name, size in (("Heading 1", 16), ("Heading 2", 14),
                       ("Heading 3", 12), ("Heading 4", 12)):
        s = doc.styles[name]
        s.font.name = FONT
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)

    sec = doc.sections[0]
    _a4(sec)
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.top_margin = sec.bottom_margin = Inches(1.0)
    _page_number_format(sec, "lowerRoman", 1)
    _add_page_field(sec.footer.paragraphs[0])

    text = re.sub(r"<!--.*?-->", "", SRC.read_text(encoding="utf-8"), flags=re.S)
    lines = text.split("\n")

    i, pending_break, current, body_started = 0, False, "", False

    while i < len(lines):
        raw, line = lines[i], lines[i].rstrip()

        if "page-break-after" in line:
            pending_break = True; i += 1; continue

        # ---- fenced code / pseudocode ------------------------------------
        if line.startswith("```"):
            j = i + 1; buf = []
            while j < len(lines) and not lines[j].startswith("```"):
                buf.append(lines[j]); j += 1
            for ln in buf:
                p = doc.add_paragraph()
                _fmt(p, spacing=1.0, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
                p.paragraph_format.left_indent = Inches(0.3)
                _shade(p, "F2F2F2")
                r = p.add_run(ln if ln.strip() else " ")
                r.font.name = MONO; r.font.size = Pt(10)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            i = j + 1; continue

        # ---- table -------------------------------------------------------
        if (line.startswith("|") and i + 1 < len(lines)
                and re.match(r"^\|[\s:|-]+\|\s*$", lines[i + 1])):
            header = [c.strip() for c in line.strip("|").split("|")]
            j = i + 2; rows = []
            while j < len(lines) and lines[j].startswith("|"):
                rows.append([c.strip() for c in lines[j].strip("|").split("|")])
                j += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for k, h in enumerate(header):
                c = t.rows[0].cells[k]; c.text = ""
                _runs(c.paragraphs[0], h, size=10, bold_all=True)
            for row in rows:
                cells = t.add_row().cells
                for k in range(len(header)):
                    cells[k].text = ""
                    _runs(cells[k].paragraphs[0], row[k] if k < len(row) else "",
                          size=10)
            for r_ in t.rows:
                for c in r_.cells:
                    for p in c.paragraphs:
                        _fmt(p, spacing=1.0, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
            i = j; continue

        # ---- image -------------------------------------------------------
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", line)
        if m:
            img = _image_for(m.group(2))
            p = doc.add_paragraph(); _fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
            if img:
                # Scale to width, but a tall diagram (the class and sequence
                # diagrams are portrait) would then run off the bottom of the
                # page, so fall back to scaling by height instead.
                from PIL import Image as _PILImage
                try:
                    with _PILImage.open(img) as im:
                        iw, ih = im.size
                    if ih / iw * MAX_IMG_IN > MAX_IMG_H_IN:
                        p.add_run().add_picture(str(img), height=Inches(MAX_IMG_H_IN))
                    else:
                        p.add_run().add_picture(str(img), width=Inches(MAX_IMG_IN))
                except Exception:
                    p.add_run().add_picture(str(img), width=Inches(MAX_IMG_IN))
            else:
                r = p.add_run("[ diagram not yet exported — save it as "
                              f"writing/figures/diagrams/{Path(m.group(2)).stem}.png ]")
                r.italic = True; r.font.name = FONT; r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)
            i += 1; continue

        # ---- heading -----------------------------------------------------
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level, title = len(m.group(1)), m.group(2).strip()
            key = title.lower().strip()
            is_front = key in FRONT
            is_chapter = bool(re.match(r"^(CHAPTER|APPENDIX)", title, re.I))

            if is_chapter and not body_started:
                body_started = True
                ns = doc.add_section(WD_SECTION.NEW_PAGE)
                _a4(ns)
                ns.left_margin = ns.right_margin = Inches(1.0)
                ns.top_margin = ns.bottom_margin = Inches(1.0)
                ns.footer.is_linked_to_previous = False
                _page_number_format(ns, "decimal", 1)
                _add_page_field(ns.footer.paragraphs[0])
                pending_break = False
            elif pending_break or is_front or is_chapter:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                pending_break = False

            if key == "front matter":
                i += 1; continue          # file scaffolding, not a document heading

            current = key

            if key in ("list of figures", "list of tables"):
                kind = "Figure" if "figure" in key else "Table"
                h = doc.add_heading(level=1)
                _runs(h, title.upper(), size=16, bold_all=True)
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _add_table_of(doc, kind)
                # skip the static markdown list that follows
                i += 1
                while i < len(lines) and not re.match(r"^#{1,6}\s", lines[i])                         and "page-break-after" not in lines[i]:
                    i += 1
                continue

            if key == "table of contents":
                h = doc.add_heading(level=1)
                _runs(h, "TABLE OF CONTENTS", size=16, bold_all=True)
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p = doc.add_paragraph()
                _fmt(p, spacing=1.5, align=WD_ALIGN_PARAGRAPH.LEFT)
                _add_toc_field(p)
                i += 1; continue

            # A level-1 heading that is not front matter is a chapter *name*
            # ("INTRODUCTION" under "CHAPTER 1") and takes chapter styling.
            big = is_chapter or is_front or level == 1
            hl = 1 if big else min(max(level, 2), 4)
            size = 16 if big else (14 if hl == 2 else 12)
            h = doc.add_heading(level=hl)
            _runs(h, title.upper() if big else title, size=size, bold_all=True)
            h.alignment = (WD_ALIGN_PARAGRAPH.CENTER if big
                           else WD_ALIGN_PARAGRAPH.LEFT)
            h.paragraph_format.space_before = Pt(12 if hl <= 2 else 8)
            h.paragraph_format.space_after = Pt(6)
            i += 1; continue

        # ---- block quote --------------------------------------------------
        if line.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].startswith(">"):
                buf.append(lines[i].lstrip("> ").rstrip()); i += 1
            p = doc.add_paragraph(); _fmt(p)
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.3)
            _shade(p, "FAF6E9")
            _runs(p, " ".join(buf))
            for r in p.runs:
                r.italic = True
            continue

        # ---- lists ---------------------------------------------------------
        m = re.match(r"^(\s*)([-*])\s+(.*)$", raw)
        if m:
            p = doc.add_paragraph(style="List Bullet" if len(m.group(1)) < 2
                                  else "List Bullet 2")
            _fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, after=3)
            _runs(p, m.group(3)); i += 1; continue
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", raw)
        if m:
            p = doc.add_paragraph(style="List Number" if len(m.group(1)) < 2
                                  else "List Number 2")
            _fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, after=3)
            _runs(p, m.group(3)); i += 1; continue

        if re.match(r"^-{3,}\s*$", line) or not line.strip():
            i += 1; continue

        # ---- paragraph (join hard-wrapped lines) ---------------------------
        buf = [line]; j = i + 1
        while (j < len(lines) and lines[j].strip()
               and not re.match(r"^(#{1,6}\s|\||```|>|\s*[-*]\s|\s*\d+\.\s|!\[)",
                                lines[j])
               and "page-break-after" not in lines[j]):
            buf.append(lines[j].rstrip()); j += 1
        body = " ".join(buf)

        cap = re.match(r"^\*\*(Figure|Table)\s+\d+\s*:\s*(.*?)\*\*\s*(.*)$", body, re.S)
        if cap:
            _add_caption(doc, cap.group(1), cap.group(2).strip())
            if cap.group(3).strip():          # analysis that followed the caption
                q = doc.add_paragraph(); _fmt(q)
                _runs(q, cap.group(3).strip())
            i = j; continue

        p = doc.add_paragraph()
        if current in DOUBLE_14:
            _fmt(p, spacing=2.0)
            _runs(p, body, size=14)
        elif current in LISTY:
            _fmt(p, spacing=1.5, align=WD_ALIGN_PARAGRAPH.LEFT)
            _runs(p, body)
        elif current.startswith("5.5") or current == "references":
            _fmt(p, spacing=1.0, align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
            _runs(p, body)
        else:
            _fmt(p)
            _runs(p, body)
        i = j

    return doc


STAMP = OUT.with_suffix(".build.json")


def _fingerprint() -> dict:
    st = OUT.stat()
    return {"size": st.st_size, "mtime": round(st.st_mtime, 3)}


def stamp_build() -> None:
    """Record the file this tool just wrote, so later edits are visible."""
    STAMP.write_text(json.dumps(_fingerprint(), indent=2) + "\n", encoding="utf-8")


def hand_formatted() -> bool:
    """Whether the .docx has been worked on since this tool last wrote it.

    This writes the whole document from the markdown, so a rebuild discards
    anything done in Word -- spacing, page breaks, table widths, the front
    matter someone typed in. None of that can live in the markdown, so the
    markdown must not silently win.

    The test is against a record of what this tool produced, not against the
    markdown's clock. Comparing the two files' times looked right and was not:
    rebuilding the markdown made it the newer file and quietly unlocked
    overwriting a document somebody had spent an afternoon formatting.

    With no record at all the file is treated as hand-formatted, because the
    costly mistake is the one that destroys work.
    """
    if not OUT.exists():
        return False
    if not STAMP.exists():
        return True
    try:
        was = json.loads(STAMP.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return True
    return _fingerprint() != was


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"{SRC} not found -- run `python -m tools.build_report` first")
    if hand_formatted() and "--force" not in sys.argv:
        raise SystemExit(
            f"  {OUT.name} has been edited since this tool wrote it, so it\n"
            "  carries formatting the markdown cannot express. Rebuilding would\n"
            "  replace the whole file and lose it." + chr(10) +
            "\n"
            "  Check first whether a rebuild is needed at all:\n"
            "      python -m tools.build_report      # names any figure that is\n"
            "                                        # missing or out of date\n"
            "\n"
            "  If it is, copy the file somewhere safe, then:\n"
            "      python -m tools.make_docx --force\n"
            "  and re-apply the Word formatting afterwards.")
    doc = build()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(OUT)
    except PermissionError:
        raise SystemExit(f"{OUT} is open in Word -- close it and re-run")

    stamp_build()
    print(f"  wrote {OUT}  ({OUT.stat().st_size/1024:,.0f} KB)")
    print(f"  {len(doc.paragraphs):,} paragraphs | {len(doc.tables)} tables "
          f"| {len(doc.inline_shapes)} images")
    print("  in Word: right-click the contents page -> Update Field")


if __name__ == "__main__":
    main()
